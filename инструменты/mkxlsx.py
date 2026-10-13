#!/usr/bin/env python3
"""Собирает delta-market.xlsx (урок 6.6), app-journal.xlsx (урок 4.7) и vilka.xlsx
(модуль 7) из data.js и data-vilka.js.

    python3 инструменты/mkxlsx.py           # записать файлы в корень сайта
    python3 инструменты/mkxlsx.py --check   # сверить с файлами в ветке

Листы orders, users — как в базе (даты — датами Excel, суммы — числами),
leads — сырые строки выгрузки CRM, для задач на чистку. В app-journal —
журнал приложения app_users, app_activity, app_orders: даты — датами
Excel, revenue — числом. Сверка идёт по
распакованному содержимому листов, а не по байтам: zlib на разных
системах сжимает по-разному.
"""
import csv, datetime, io, pathlib, re, sys, zipfile
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = pathlib.Path(__file__).resolve().parent.parent
STAMP = datetime.datetime(2024, 9, 30)
SRC = (ROOT / "data.js").read_text(encoding="utf-8")
VILKA = (ROOT / "data-vilka.js").read_text(encoding="utf-8")


def rows(key, src):
    body = re.search(r"\n  %s: `(.*?)`" % key, src, re.S).group(1)
    return [r for r in csv.reader(io.StringIO(body)) if r]


def day(s):
    return datetime.datetime.strptime(s, "%Y-%m-%d")


def stamp(s):
    return datetime.datetime.strptime(s, "%Y-%m-%d %H:%M")


BOOKS = {
    "delta-market.xlsx": (
        ("orders", "ordersCSV", (int, int, day, float, str), ((3, "yyyy-mm-dd"), (4, "0.00"))),
        ("users", "usersCSV", (int, day, str, str, str), ((2, "yyyy-mm-dd"),)),
        ("leads", "leadsCSV", None, ()),   # как есть: текст, пустые — пустые
    ),
    "app-journal.xlsx": (
        ("app_users", "appUsersCSV", (int, day, str, str, str), ((2, "yyyy-mm-dd"),)),
        ("app_activity", "appActivityCSV", (int, day), ((2, "yyyy-mm-dd"),)),
        ("app_orders", "appOrdersCSV", (int, day, float), ((2, "yyyy-mm-dd"), (3, "0.00"))),
    ),
    "vilka.xlsx": (
        ("customers", "customersCSV", (int, day, str, str, str), ((2, "yyyy-mm-dd"),)),
        ("restaurants", "restaurantsCSV", (int, str, str), ()),
        ("orders", "ordersCSV", (int, int, int, stamp, int, int, str), ((4, "yyyy-mm-dd hh:mm"),)),
        ("experiment", "experimentCSV", (int, str, day), ((3, "yyyy-mm-dd"),)),
    ),
}


def build(sheets, src):
    wb = Workbook()
    wb.properties.creator = "Тетрадь аналитика"
    wb.properties.created = wb.properties.modified = STAMP
    wb.remove(wb.active)
    for name, key, cast, fmts in sheets:
        data = rows(key, src)
        ws = wb.create_sheet(name)
        ws.append(data[0])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in data[1:]:
            ws.append([f(v) for f, v in zip(cast, r)] if cast else [v if v != "" else None for v in r])
        for col, fmt in fmts:
            for (c,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                c.number_format = fmt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    raw = io.BytesIO()
    wb.save(raw)
    # одинаковые даты у файлов внутри zip — иначе каждая сборка «новая»
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            z = zipfile.ZipInfo(info.filename, date_time=STAMP.timetuple()[:6])
            z.compress_type = zipfile.ZIP_DEFLATED
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl при сохранении ставит modified = сейчас — возвращаем STAMP
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + STAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode(), data)
            dst.writestr(z, data)
    return out.getvalue()


def parts(blob):
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        return {n: z.read(n) for n in z.namelist()}


if __name__ == "__main__":
    stale = []
    for name, sheets in BOOKS.items():
        # у «Вилки» свои orders и customers — ключи совпадают с data.js, источник задаёт книга
        out, blob = ROOT / name, build(sheets, VILKA if name == "vilka.xlsx" else SRC)
        if "--check" in sys.argv:
            if not out.exists() or parts(out.read_bytes()) != parts(blob):
                stale.append(name)
            else:
                print("%s совпадает с данными" % name)
        else:
            out.write_bytes(blob)
            print("Готово: %s — %.0f КБ" % (name, out.stat().st_size / 1024))
    if stale:
        sys.exit("%s устарел: python3 инструменты/mkxlsx.py" % ", ".join(stale))
